import streamlit as st
import re
from docx import Document
from docxtpl import DocxTemplate
import PyPDF2
import io

st.set_page_config(page_title="AI CV Mapper", page_icon="📄", layout="wide")

def extract_text_from_pdf(pdf_file):
    """Extract text from PDF file"""
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(docx_file):
    """Extract text from DOCX file"""
    doc = Document(docx_file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def parse_cv_with_rules(cv_text):
    """Use enhanced rule-based parsing to extract CV data (FREE - no API needed)"""
    data = {
        'name': '',
        'email': '',
        'phone': '',
        'address': '',
        'summary': '',
        'education': [],
        'experience': [],
        'skills': [],
        'certifications': [],
        'projects': []
    }
    
    lines = [line.strip() for line in cv_text.split('\n') if line.strip()]
    text_lower = cv_text.lower()
    
    # Extract email with multiple patterns
    email_patterns = [
        r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}',
        r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    ]
    for pattern in email_patterns:
        email_matches = re.findall(pattern, cv_text)
        if email_matches:
            data['email'] = email_matches[0]
            break
    
    # Extract phone with enhanced patterns
    phone_patterns = [
        r'\+?\d{1,3}[\s.-]?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}',
        r'\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}',
        r'\d{3}[\s.-]\d{3}[\s.-]\d{4}',
        r'\d{10}',
        r'\+\d{11,}'
    ]
    for pattern in phone_patterns:
        phone_matches = re.findall(pattern, cv_text)
        if phone_matches:
            data['phone'] = phone_matches[0]
            break
    
    # Extract name (improved logic)
    # Look for name patterns in first few lines
    for i, line in enumerate(lines[:8]):
        line_clean = line.strip()
        # Skip lines with email or phone
        if any(char in line_clean for char in ['@', '(', ')', '+']):
            continue
        # Look for lines that could be names
        if (len(line_clean.split()) <= 4 and 
            len(line_clean) > 5 and 
            len(line_clean) < 50 and
            not any(keyword in line_clean.lower() for keyword in ['address', 'phone', 'email', 'cv', 'resume', 'curriculum'])):
            # Check if it looks like a name (mostly letters)
            if re.match(r'^[A-Za-z\s.,-]+$', line_clean):
                data['name'] = line_clean
                break
    
    # Extract sections with improved logic
    current_section = None
    section_content = {}
    
    # Define section keywords with variations
    section_keywords = {
        'summary': ['summary', 'objective', 'profile', 'about', 'overview', 'introduction'],
        'education': ['education', 'academic', 'qualification', 'degree', 'university', 'school', 'college'],
        'experience': ['experience', 'employment', 'work', 'career', 'professional', 'job', 'position'],
        'skills': ['skills', 'technical', 'competencies', 'expertise', 'technologies', 'abilities'],
        'certifications': ['certification', 'certificates', 'licenses', 'awards', 'achievements'],
        'projects': ['projects', 'portfolio', 'work samples', 'personal projects']
    }
    
    i = 0
    while i < len(lines):
        line = lines[i]
        line_lower = line.lower().strip()
        
        # Check if this line is a section header
        detected_section = None
        for section, keywords in section_keywords.items():
            for keyword in keywords:
                if (keyword in line_lower and 
                    len(line.split()) <= 5 and
                    len(line) < 100):
                    detected_section = section
                    break
            if detected_section:
                break
        
        if detected_section:
            current_section = detected_section
            section_content[current_section] = []
            i += 1
            continue
        
        # Add content to current section
        if current_section and line.strip():
            # Skip very short lines and lines that look like headers
            if len(line.strip()) > 5 and not line.strip().endswith(':'):
                section_content[current_section].append(line.strip())
        
        i += 1
    
    # Process sections and clean up
    for section, content in section_content.items():
        if section == 'summary':
            # Join summary lines into a paragraph
            data['summary'] = ' '.join(content[:3])  # Take first 3 lines for summary
        elif section in ['education', 'experience', 'skills', 'certifications', 'projects']:
            # Clean and filter content
            cleaned_content = []
            for item in content:
                if len(item) > 10 and len(item) < 500:  # Reasonable length
                    cleaned_content.append(item)
            data[section] = cleaned_content[:10]  # Limit to 10 items per section
    
    # Fallback: if no sections found, try to extract from raw text
    if not any(data[key] for key in ['education', 'experience', 'skills']):
        # Simple fallback extraction
        sentences = [s.strip() for s in cv_text.split('.') if len(s.strip()) > 20]
        data['summary'] = sentences[0] if sentences else 'Professional with relevant experience'
        
        # Look for common patterns
        for sentence in sentences[:20]:
            sentence_lower = sentence.lower()
            if any(word in sentence_lower for word in ['university', 'degree', 'graduated', 'bachelor', 'master']):
                data['education'].append(sentence[:200])
            elif any(word in sentence_lower for word in ['worked', 'employed', 'company', 'role', 'position']):
                data['experience'].append(sentence[:200])
            elif any(word in sentence_lower for word in ['skill', 'proficient', 'experience with', 'knowledge']):
                data['skills'].append(sentence[:200])
    
    # Ensure all fields have some content
    if not data['name'] and data['email']:
        data['name'] = data['email'].split('@')[0].replace('.', ' ').title()
    
    if not data['summary']:
        data['summary'] = 'Experienced professional with a strong background in their field.'
    
    return data

def fill_template(template_file, data):
    """Fill DOCX template with extracted data"""
    doc = DocxTemplate(template_file)
    
    # Convert lists to formatted strings
    context = {}
    for key, value in data.items():
        if isinstance(value, list) and value:
            # Format list items with bullet points
            context[key] = "\n".join([f"• {item}" for item in value])
        elif isinstance(value, list):
            context[key] = "No information available"
        else:
            context[key] = value if value else "Not specified"
    
    doc.render(context)
    
    # Save to bytes
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def main():
    st.title("🤖 AI CV Mapper Agent (FREE Version)")
    st.markdown("**Automatically extract content from your old CV and populate a new template**")
    st.info("✨ This version uses enhanced smart parsing - NO API KEY NEEDED!")
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.header("📤 Upload Old CV")
        old_cv = st.file_uploader("Upload your old CV (PDF or DOCX)", type=["pdf", "docx"], key="old")
        
        if old_cv:
            st.success(f"✅ Uploaded: {old_cv.name}")
    
    with col2:
        st.header("📥 Upload New Template")
        new_template = st.file_uploader("Upload new CV template (DOCX with placeholders)", type=["docx"], key="new")
        
        if new_template:
            st.success(f"✅ Uploaded: {new_template.name}")
            with st.expander("ℹ️ Template Instructions"):
                st.markdown("""
                Your DOCX template should contain placeholders like:
                - `{{name}}` - Full name
                - `{{email}}` - Email address
                - `{{phone}}` - Phone number
                - `{{address}}` - Address
                - `{{summary}}` - Professional summary
                - `{{education}}` - Education details
                - `{{experience}}` - Work experience
                - `{{skills}}` - Skills list
                - `{{certifications}}` - Certifications
                - `{{projects}}` - Projects
                """)
    
    if st.button("🚀 Generate New CV", type="primary", use_container_width=True):
        if not old_cv or not new_template:
            st.error("❌ Please upload both old CV and new template!")
            return
        
        with st.spinner("🔄 Processing your CV..."):
            # Extract text from old CV
            if old_cv.name.endswith('.pdf'):
                cv_text = extract_text_from_pdf(old_cv)
            else:
                cv_text = extract_text_from_docx(old_cv)
            
            st.info("✓ Extracted text from old CV")
            
            # Parse with enhanced rules
            try:
                parsed_data = parse_cv_with_rules(cv_text)
                st.info("✓ Parsed CV with enhanced intelligent rules")
                
                # Show extracted data
                with st.expander("📋 View Extracted Data"):
                    st.json(parsed_data)
                
                # Fill template
                filled_cv = fill_template(new_template, parsed_data)
                st.success("✅ Successfully generated new CV!")
                
                # Download button
                st.download_button(
                    label="⬇️ Download New CV",
                    data=filled_cv,
                    file_name="new_cv_filled.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary",
                    use_container_width=True
                )
                
            except Exception as e:
                st.error(f"❌ Error: {str(e)}")
                st.exception(e)

if __name__ == "__main__":
    main()
